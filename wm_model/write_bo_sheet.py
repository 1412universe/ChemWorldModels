"""从 _bo_result.json（世界模型BO）渲染「BO推荐」sheet：Top4 + 24点推演。"""
from pathlib import Path
_ROOT = Path(__file__).resolve().parent.parent
_DATA = _ROOT / "data"
_WM = Path(__file__).resolve().parent
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

XLSX = str(_DATA / "OCM_master.xlsx")
R = json.load(open(_WM / "_bo_result.json"))
chosen, points = R["chosen"], R["points"]

REASON = {
    ("Pd", "ZnO(70±30nm)"): "实测n=3最高产率档(sel~95%)；Pd强CH3结合、ZnO CB≈-0.35V供还原电子，世界模型act/red协同，首选复现",
    ("Ag", "P25(TiO2)"): "实测n=8最稳健基线；Ag低功函数利电子流出、P25标准UV载体，作对照体系",
    ("Au", "CeO2"): "实测n=1；Au温和活化O2、CeO2储氧/氧空位调O*，待复现",
    ("Cu", "SrTiO3(100nm)"): "实测n=1；Cu中等en/wf、SrTiO3 CB较负供还原、化学稳定，待复现提量",
}


def win(n):
    a, b = (n - 1) * 18, n * 18
    return f"{a//60:02d}:{a%60:02d}–{b//60:02d}:{b%60:02d}"


wb = openpyxl.load_workbook(XLSX)
if "BO推荐" in wb.sheetnames:
    del wb["BO推荐"]
ws = wb.create_sheet("BO推荐")

navy = "3370FF"
thin = Side(style="thin", color="B8C7D9")
border = Border(bottom=thin)
hfont = Font(name="微软雅黑", bold=True, color="FFFFFF", size=10)
sfont = Font(name="微软雅黑", bold=True, size=10, color="133C9A")
bfont = Font(name="微软雅黑", size=10)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left = Alignment(horizontal="left", vertical="center", wrap_text=True)

r = 1
ws.cell(r, 1, "世界模型BO：4催化剂×各自6组=24点；E1为实测/模型锚点，E2-E6为世界模型条件推演；每组18min，温度/光强锁，单变量步进").font = Font(name="微软雅黑", bold=True, size=11)
r += 2

# 块A Top4
_n_obs = sum(1 for d in chosen if d.get("n_obs", 0) and d["n_obs"] > 0)
_tail = "全部有实测锚点" if _n_obs == len(chosen) else f"{_n_obs}实测锚点，其余为世界模型探索位"
ws.cell(r, 1, f"一、Top4（{_tail}；UCB=μ_y+0.55·sel/100+探索项）").font = sfont
r += 1
headA = ["排名", "催化剂", "金属", "载体", "en", "wf/eV", "Eg", "CB/V", "VB/V",
         "锚点产率/μmol·g⁻¹·h⁻¹", "锚点sel/%", "证据", "世界模型act", "red", "理由"]
for j, h in enumerate(headA, 1):
    c = ws.cell(r, j, h); c.fill = PatternFill("solid", fgColor=navy); c.font = hfont; c.alignment = center; c.border = border
r += 1
for i, d in enumerate(chosen, 1):
    reason = REASON.get((d["metal"], d["sup"]), "世界模型/实测组合，建议复现")
    vals = [i, f'{d["metal"]}/{d["sup"]}', d["metal"], d["sup"], d["en"], d["wf"],
            d["Eg"], d["CB"], d["VB"], round(10 ** d["mu_y"]), round(d["mu_s"], 1),
            d["tag"], "", "", reason]
    # act/red 取该催化剂基线点
    bp = next(p for p in points if p["metal"] == d["metal"] and p["sup"] == d["sup"] and p["inner"] == "E1")
    vals[12], vals[13] = bp["act"], bp["red"]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = bfont; c.alignment = left if j == 15 else center
    r += 1

r += 1
# 块B 24点
ws.cell(r, 1, "二、24点（世界模型推演预测y/sel；E1锚点，其余为条件响应估算，待实验验证）").font = sfont
r += 1
headB = ["实验#", "催化剂", "组内", "步进", "CH4", "O2", "Ar", "总流量",
        "时间窗", "预测y", "预测sel/%", "flow_sel", "act", "red", "备注"]
for j, h in enumerate(headB, 1):
    c = ws.cell(r, j, h); c.fill = PatternFill("solid", fgColor=navy); c.font = hfont; c.alignment = center; c.border = border
r += 1
for p in points:
    note = "锚点(E1)" if p["inner"] == "E1" else "世界模型条件估算"
    vals = [p["exp"], f'{p["metal"]}/{p["sup"]}', p["inner"], p["step"], p["ch4"], p["o2"],
            p["ar"], p["tot"], win(p["exp"]), p["pred_y"], p["pred_sel"],
            p["flow_sel"], p["act"], p["red"], note]
    for j, v in enumerate(vals, 1):
        c = ws.cell(r, j, v); c.font = bfont; c.alignment = left if j == 15 else center
    r += 1

widths = [6, 18, 6, 7, 6, 6, 6, 7, 14, 9, 9, 8, 7, 7, 20]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(j)].width = w
wb.save(XLSX)
print("written; points =", len(points))
