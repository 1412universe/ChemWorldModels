# -*- coding: utf-8 -*-
"""工作流每次运行的数据日志：整理成表格，xlsx 写到 data/Datalog_Workflow。"""
from __future__ import annotations

import json
import os
import re
import threading
from datetime import datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HERE = os.path.dirname(os.path.abspath(__file__))
_REPO = os.path.dirname(os.path.dirname(_HERE))
DATALOG_DIR = os.path.join(_REPO, "data", "Datalog_Workflow")
_INDEX = os.path.join(DATALOG_DIR, "index.json")
_LOCK = threading.Lock()

COLUMNS = [
    "工作流", "运行ID", "模式", "节点", "设备", "动作", "状态",
    "开始时间", "结束时间", "耗时_秒",
    "实际称量_g", "实际移液_uL",
    "烘干时间", "烘干温度",
    "离心时间", "震荡时间", "超声时间",
    "吸放方向", "枪头位", "吸液板位", "放液板位",
    "取值来源", "设定与结果",
]


def _num(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _flat(info: Any) -> Dict[str, Any]:
    out: Dict[str, Any] = {}
    if not isinstance(info, dict):
        return out
    out.update(info)
    nested = info.get("result")
    if isinstance(nested, dict):
        for k, v in nested.items():
            if k not in out or out[k] in (None, ""):
                out[k] = v
    return out


def _param(node: Dict[str, Any]) -> Dict[str, Any]:
    p = node.get("param")
    return dict(p) if isinstance(p, dict) else {}


def _pick(src: Dict[str, Any], keys: List[str]) -> Any:
    for k in keys:
        if k in src and src[k] not in (None, ""):
            return src[k]
    return None


def _fmt_time(ts: Optional[float]) -> str:
    if not ts:
        return ""
    try:
        return datetime.fromtimestamp(float(ts)).strftime("%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError, OSError):
        return ""


def _safe(name: str) -> str:
    s = re.sub(r"[^\w\u4e00-\u9fff.\-]+", "_", (name or "workflow").strip())
    return (s or "workflow")[:40]


def _measure(node: Dict[str, Any], info: Dict[str, Any], state: str) -> Dict[str, str]:
    """从节点参数和驱动返回里抽出称量、移液、烘干、离心、震荡、超声。"""
    p = _param(node)
    f = _flat(info)
    dev = str(node.get("device_name") or "")
    name = str(node.get("name") or node.get("template_name") or "")
    label = str(node.get("footer") or node.get("_label") or name)
    blob = " ".join([dev, name, label, str(p.get("cmd_type") or "")]).lower()
    row = {c: "" for c in COLUMNS[10:]}
    sources = []

    def put(col: str, val: Any, source: str) -> None:
        if val is None or val == "":
            return
        row[col] = val if isinstance(val, str) else ("%g" % val if isinstance(val, float) and val == int(val) else val)
        if isinstance(val, float):
            row[col] = ("%g" % val) if abs(val - round(val)) < 1e-6 else ("%.4g" % val)
        if source not in sources:
            sources.append(source)

    # 配平节点：两块板的实测重量 + 交给 9320 的吸放液
    if f.get("reaction_weight_g") is not None or f.get("counterweight_weight_g") is not None:
        w1 = _num(f.get("reaction_weight_g"))
        w2 = _num(f.get("counterweight_weight_g"))
        parts = []
        if w1 is not None:
            parts.append("反应板 %.2f" % w1)
        if w2 is not None:
            parts.append("配平板 %.2f" % w2)
        put("实际称量_g", "；".join(parts), "实测")
        put("实际移液_uL", f.get("volume_ul"), "实测")
        put("吸放方向", f.get("direction"), "实测")
        put("枪头位", f.get("tip_slot") and ("T%s" % f.get("tip_slot")), "实测")
        put("吸液板位", f.get("source_slot") and ("T%s" % f.get("source_slot")), "实测")
        put("放液板位", f.get("target_slot") and ("T%s" % f.get("target_slot")), "实测")

    weight = _pick(f, ["actual_weight_g", "weight_g", "weight", "mass_g", "net_g"])
    if weight is None:
        weight = _pick(p, ["weight", "mass", "target_weight", "weight_g"])
        if weight is not None and ("weigh" in blob or "称" in label or dev == "gn_solid_weighing"):
            put("实际称量_g", weight, "下发参数" if state != "dry" else "干跑")
    else:
        put("实际称量_g", weight, "实测")

    vol = _pick(f, ["volume_ul", "delta_ul", "delta_ul_water", "actual_volume_ul"])
    if vol is None:
        vol = _pick(p, ["volume", "volume_ul", "aspirate_volume", "dispense_volume", "ul"])
        if vol is not None and ("prcxi" in blob or "liquid" in blob or "移液" in label or "aspirat" in blob or "dispens" in blob):
            put("实际移液_uL", vol, "干跑" if state == "dry" else "下发参数")
    elif "实际移液_uL" not in row or not row["实际移液_uL"]:
        put("实际移液_uL", vol, "实测")

    cmd = str(p.get("cmd_type") or "")
    dry = state == "dry"
    is_oven = dev in ("gn_standard_oven", "gn_vacuum_oven") or "烘干" in label or "oven" in blob
    is_centrifuge = dev == "gn_centrifuge" and name != "balance_plates" and (cmd == "6" or name == "run")
    is_shake = any(k in blob or k in label for k in ("shake", "shaker", "震荡", "振荡"))
    is_sonic = any(k in blob or k in label for k in ("ultra", "sonic", "超声"))

    def origin(measured_key: str) -> str:
        if f.get(measured_key) is not None:
            return "实测"
        return "干跑" if dry else "下发参数"

    if is_oven:
        put("烘干时间", _pick(f, ["actual_dry_time"]) or _pick(p, ["time_minutes", "time", "minutes"]), origin("actual_dry_time"))
        put("烘干温度", _pick(f, ["actual_temperature", "temperature_c"]) or _pick(p, ["temperature", "temp"]), origin("actual_temperature"))
    if is_centrifuge:
        put("离心时间", _pick(f, ["actual_centrifuge_time"]) or _pick(p, ["time_minutes", "minutes", "time"]), origin("actual_centrifuge_time"))
    if is_shake:
        put("震荡时间", _pick(f, ["actual_shake_time"]) or _pick(p, ["shake_time", "mix_time", "time_minutes", "time", "minutes"]), origin("actual_shake_time"))
    if is_sonic:
        put("超声时间", _pick(f, ["actual_ultrasound_time"]) or _pick(p, ["ultrasound_time", "sonic_time", "time", "seconds"]), origin("actual_ultrasound_time"))

    bits = []
    for k in ("cmd_type", "time_minutes", "temperature", "rpm", "volume", "weight"):
        if p.get(k) not in (None, ""):
            bits.append("%s=%s" % (k, p.get(k)))
    msg = f.get("message") or f.get("error") or ""
    if msg:
        bits.append(str(msg))
    row["取值来源"] = "、".join(sources)
    row["设定与结果"] = "；".join(bits)
    return row


def build_rows(run: Any) -> List[Dict[str, Any]]:
    nodes = getattr(run, "nodes", {}) or {}
    order = list(getattr(run, "order", []) or [])
    for u, n in nodes.items():
        if u not in order:
            order.append(u)
    states = getattr(run, "states", {}) or {}
    timing = getattr(run, "timing", {}) or {}
    info = getattr(run, "info", {}) or {}
    rows = []
    for u in order:
        n = nodes.get(u) or {}
        if n.get("type") == "Group" or n.get("name") == "group":
            continue
        t = timing.get(u) or {}
        st = states.get(u) or ""
        measured = _measure(n, info.get(u) or {}, "dry" if getattr(run, "mode", "") == "dry" else st)
        label = str(n.get("footer") or n.get("_label") or n.get("name") or "")
        rows.append({
            "工作流": getattr(run, "workflow_name", "") or "",
            "运行ID": getattr(run, "run_id", "") or "",
            "模式": getattr(run, "mode", "") or "",
            "节点": label or n.get("name") or u,
            "设备": n.get("device_name") or "",
            "动作": n.get("name") or "",
            "状态": st,
            "开始时间": _fmt_time(t.get("start_wall")),
            "结束时间": _fmt_time(t.get("end_wall")),
            "耗时_秒": t.get("dur") if t.get("dur") is not None else "",
            **measured,
        })
    return rows


def _load_index() -> List[Dict[str, Any]]:
    if not os.path.isfile(_INDEX):
        return []
    try:
        with open(_INDEX, encoding="utf-8") as f:
            data = json.load(f)
        return list(data.get("runs") or [])
    except (OSError, json.JSONDecodeError):
        return []


def _write_xlsx(path: str, rows: List[Dict[str, Any]]) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "运行日志"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="334155")
    for col, name in enumerate(COLUMNS, 1):
        cell = ws.cell(1, col, name)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, 2):
        for c, name in enumerate(COLUMNS, 1):
            ws.cell(r, c, row.get(name, ""))
    for i, name in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(36, max(12, len(name) + 4))
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLUMNS)), max(1, len(rows) + 1))
    ws.freeze_panes = "A2"
    wb.save(path)


def save_workflow_run(run: Any) -> Dict[str, Any]:
    """把一次运行写成 xlsx，并记进 index.json。同一 run_id 再写会覆盖。"""
    os.makedirs(DATALOG_DIR, exist_ok=True)
    rows = build_rows(run)
    run_id = str(getattr(run, "run_id", "") or "")
    name = _safe(getattr(run, "workflow_name", "") or "workflow")
    fname = "%s_%s.xlsx" % (name, run_id or "run")
    path = os.path.join(DATALOG_DIR, fname)
    _write_xlsx(path, rows)
    entry = {
        "run_id": run_id,
        "workflow_name": getattr(run, "workflow_name", "") or "",
        "workflow_uuid": getattr(run, "workflow_uuid", "") or "",
        "mode": getattr(run, "mode", "") or "",
        "started_at": _fmt_time(getattr(run, "started_at", 0)),
        "ended_at": _fmt_time(getattr(run, "ended_at", 0)),
        "finished": bool(getattr(run, "finished", False)),
        "rows": len(rows),
        "xlsx": fname,
    }
    with _LOCK:
        runs = [r for r in _load_index() if r.get("run_id") != run_id]
        runs.insert(0, entry)
        tmp = _INDEX + ".tmp"
        with open(tmp, "w", encoding="utf-8") as f:
            json.dump({"runs": runs[:400]}, f, ensure_ascii=False, indent=2)
        os.replace(tmp, _INDEX)
    return entry


def list_runs() -> List[Dict[str, Any]]:
    with _LOCK:
        return _load_index()


def run_detail(run_id: str) -> Optional[Dict[str, Any]]:
    rid = str(run_id or "")
    hit = None
    for r in list_runs():
        if r.get("run_id") == rid:
            hit = dict(r)
            break
    if not hit:
        return None
    path = os.path.join(DATALOG_DIR, hit.get("xlsx") or "")
    rows: List[Dict[str, Any]] = []
    if os.path.isfile(path):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for raw in ws.iter_rows(min_row=2, values_only=True):
            if not any(v not in (None, "") for v in raw):
                continue
            rows.append({str(header[i] or ""): ("" if raw[i] is None else raw[i]) for i in range(len(header))})
        wb.close()
    hit["records"] = rows
    return hit


def xlsx_path(run_id: str) -> Optional[str]:
    rid = str(run_id or "")
    for r in list_runs():
        if r.get("run_id") == rid:
            path = os.path.join(DATALOG_DIR, r.get("xlsx") or "")
            return path if os.path.isfile(path) else None
    return None
